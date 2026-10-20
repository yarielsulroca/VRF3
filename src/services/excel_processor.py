import pandas as pd
from pathlib import Path
from typing import List, Dict, Optional, Callable, Tuple
import threading
import time
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler

from src.config import (
    COLUMN_MAP,
    REQUIRED_COLUMNS,
    APP_CONFIG,
    HEADER_SCAN_MAX_ROWS,
    column_label,
    find_column,
    LIBRO_SCHEMA,
    match_schema_column,
)
from src.models.invoice import Invoice, TaxSummary


def _find_header_row(path: Path) -> int:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        for row_idx, row in enumerate(
            worksheet.iter_rows(max_row=HEADER_SCAN_MAX_ROWS, values_only=True)
        ):
            labels = [column_label(cell).lower() for cell in row if cell is not None]
            if not labels:
                continue
            has_fecha = any("fecha" in label for label in labels)
            has_total = any(label == "total" or label.startswith("total") for label in labels)
            has_numero = any(label in ("numero", "número", "nro") or "numero" in label for label in labels)
            if has_fecha and has_total and has_numero:
                return row_idx
    finally:
        workbook.close()
    return 0


def _estimate_data_row_count(path: Path, header_row: int) -> int:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        max_row = worksheet.max_row or 1
    finally:
        workbook.close()
    return max(0, max_row - (header_row + 1))


def _read_excel_chunked(
    path: Path,
    chunk_size: int,
    header_row: int = 0,
    progress_callback: Optional[Callable[[int, int], None]] = None,
) -> pd.DataFrame:
    total_estimate = _estimate_data_row_count(path, header_row)
    if progress_callback:
        progress_callback(0, max(total_estimate, 1))

    chunks: List[pd.DataFrame] = []
    skip_data_rows = 0
    columns: Optional[List] = None
    total_rows = 0

    while True:
        if skip_data_rows == 0:
            chunk = pd.read_excel(
                path,
                engine="openpyxl",
                header=header_row,
                nrows=chunk_size,
            )
            if chunk.empty:
                break
            columns = chunk.columns.tolist()
        else:
            chunk = pd.read_excel(
                path,
                engine="openpyxl",
                skiprows=range(0, header_row + 1 + skip_data_rows),
                nrows=chunk_size,
                header=None,
                names=columns,
            )

        if chunk.empty:
            break

        chunks.append(chunk)
        total_rows += len(chunk)
        skip_data_rows += len(chunk)

        if progress_callback:
            progress_callback(total_rows, max(total_estimate, total_rows))

        if len(chunk) < chunk_size:
            break

    return pd.concat(chunks, ignore_index=True) if chunks else pd.DataFrame()


class ExcelProcessor:
    def __init__(self):
        self._df: Optional[pd.DataFrame] = None
        self._invoices: List[Invoice] = []
        self._col_map: Dict[str, str] = {}
        self._last_mtime: float = 0
        self._observer: Optional[Observer] = None
        self._reload_callback: Optional[Callable] = None
        self._path: Optional[Path] = None
        self.missing_money_columns: List[str] = []

    @property
    def path(self) -> Optional[Path]:
        return self._path

    def load(
        self,
        path: Path,
        progress_callback: Optional[Callable[[int, int], None]] = None,
    ) -> Tuple[List[Invoice], Dict[str, str]]:
        path = Path(path)
        if not path.exists():
            raise FileNotFoundError(f"No se encuentra el archivo: {path}")

        self._path = path
        self._last_mtime = path.stat().st_mtime

        header_row = _find_header_row(path)
        self._df = _read_excel_chunked(
            path,
            APP_CONFIG["chunk_size"],
            header_row,
            progress_callback,
        )
        self._df = self._drop_empty_rows(self._df)

        self._col_map = self._detect_columns(self._df.columns.tolist())
        self._validate_required_columns()

        self._invoices = self._parse_invoices(progress_callback)

        return self._invoices, self._col_map

    def _drop_empty_rows(self, df: pd.DataFrame) -> pd.DataFrame:
        if df.empty:
            return df
        fecha_col = find_column(df.columns.tolist(), "fecha")
        if fecha_col is None:
            return df
        cleaned = df[df[fecha_col].notna()].copy()
        return cleaned.reset_index(drop=True)

    def _detect_columns(self, df_columns: List) -> Dict[str, str]:
        col_map: Dict[str, str] = {}
        used: set = set()
        for item in LIBRO_SCHEMA:
            col = match_schema_column(df_columns, item, used)
            if col is not None:
                col_map[item["key"]] = col
                used.add(col)

        for target in COLUMN_MAP:
            if target in col_map:
                continue
            col = find_column([c for c in df_columns if c not in used], target)
            if col is not None:
                col_map[target] = col
                used.add(col)

        self.missing_money_columns = [
            item["label"] for item in LIBRO_SCHEMA if item.get("sum") and item["key"] not in col_map
        ]
        return col_map

    def _validate_required_columns(self):
        missing = []
        for req in REQUIRED_COLUMNS:
            if req not in self._col_map:
                missing.append(req)
        if missing:
            labels = {item["key"]: item["label"] for item in LIBRO_SCHEMA}
            faltan = ", ".join(labels.get(m, m) for m in missing)
            raise ValueError(
                f"No se encontraron las columnas obligatorias: {faltan}.\n"
                "Revisá que el Excel tenga la fila de encabezados del Libro IVA Compras."
            )

    def _parse_invoices(self, progress_callback: Optional[Callable[[int, int], None]] = None) -> List[Invoice]:
        invoices = []
        total = len(self._df)
        for idx, row in self._df.iterrows():
            try:
                inv = Invoice.from_series(row, self._col_map)
                if inv.fecha is None and not inv.numero_comprobante and inv.total == 0:
                    continue
                invoices.append(inv)
            except Exception as e:
                print(f"Error parseando fila {idx}: {e}")
            if progress_callback and idx % 100 == 0:
                progress_callback(idx + 1, total)
        if progress_callback:
            progress_callback(total, total)
        return invoices

    def get_invoices(self) -> List[Invoice]:
        return self._invoices

    def get_dataframe(self) -> Optional[pd.DataFrame]:
        return self._df

    def get_column_map(self) -> Dict[str, str]:
        return self._col_map.copy()

    def start_watcher(self, callback: Callable):
        """El callback se ejecuta en el hilo de watchdog, no en el de la UI."""
        self.stop_watcher()
        if self._path is None:
            return
        self._reload_callback = callback
        watched = self._path.resolve()

        class ExcelHandler(FileSystemEventHandler):
            def __init__(self, processor):
                self.processor = processor

            def on_modified(self, event):
                if Path(event.src_path).resolve() != watched:
                    return
                time.sleep(0.5)
                try:
                    current_mtime = watched.stat().st_mtime
                except OSError:
                    return
                if current_mtime != self.processor._last_mtime:
                    self.processor._last_mtime = current_mtime
                    if self.processor._reload_callback:
                        self.processor._reload_callback()

        self._observer = Observer()
        self._observer.schedule(ExcelHandler(self), str(watched.parent), recursive=False)
        self._observer.start()

    def stop_watcher(self):
        if self._observer:
            self._observer.stop()
            self._observer.join()
            self._observer = None


class TaxCalculator:
    """Sumatoria de importes del libro para visualización (sin recalcular impuestos)."""

    @staticmethod
    def calculate(invoices: List[Invoice]) -> TaxSummary:
        return TaxSummary.from_invoices(invoices)

    @staticmethod
    def calculate_by_tipo(invoices: List[Invoice]) -> Dict[str, TaxSummary]:
        groups: Dict[str, List[Invoice]] = {}
        for inv in invoices:
            groups.setdefault(inv.tipo_comprobante or "SIN TIPO", []).append(inv)
        return {tipo: TaxSummary.from_invoices(items) for tipo, items in groups.items()}

    @staticmethod
    def calculate_by_date(invoices: List[Invoice]) -> Dict[str, TaxSummary]:
        groups: Dict[str, List[Invoice]] = {}
        for inv in invoices:
            key = inv.fecha.strftime("%Y-%m") if inv.fecha else "SIN FECHA"
            groups.setdefault(key, []).append(inv)
        return {key: TaxSummary.from_invoices(groups[key]) for key in sorted(groups)}